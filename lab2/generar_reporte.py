from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Estilos globales ──────────────────────────────────────────────────────────
style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(11)

# Márgenes
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

# ── Colores ───────────────────────────────────────────────────────────────────
AZUL_TITULO   = RGBColor(0x00, 0x47, 0xAB)   # azul cobalto
AZUL_SUB      = RGBColor(0x1F, 0x49, 0x7D)
GRIS_CODIGO   = RGBColor(0x2B, 0x2B, 0x2B)
BG_CODIGO_HEX = "F3F3F3"                       # fondo gris claro para bloques de código

# ── Helpers ───────────────────────────────────────────────────────────────────

def add_heading(doc, text, level=1, color=AZUL_TITULO):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = color
    if level == 1:
        run.font.size = Pt(18)
    elif level == 2:
        run.font.size = Pt(14)
    else:
        run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def shade_cell(cell, fill_hex):
    """Aplica color de fondo a una celda de tabla."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill_hex)
    tcPr.append(shd)


def add_code_block(doc, code_text, label=None):
    """Agrega un bloque de código con fondo gris y fuente monoespaciada."""
    if label:
        lp = doc.add_paragraph()
        lr = lp.add_run(f"  {label}")
        lr.bold = True
        lr.font.size = Pt(9)
        lr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        lp.paragraph_format.space_after = Pt(0)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = tbl.cell(0, 0)
    shade_cell(cell, BG_CODIGO_HEX)
    cell.width = Inches(6.2)

    # Borde sutil
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'),   'single')
        b.set(qn('w:sz'),    '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'CCCCCC')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.clear()
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = GRIS_CODIGO
    p.paragraph_format.left_indent  = Pt(6)
    p.paragraph_format.right_indent = Pt(6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image_placeholder(doc, label="[Insertar captura de pantalla aquí]"):
    """Agrega un recuadro visual como marcador de imagen."""
    tbl  = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    shade_cell(cell, "EAF2FB")
    cell.width = Inches(5.5)

    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'),   'dashed')
        b.set(qn('w:sz'),    '6')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '4472C4')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"\n🖼  {label}\n")
    run.font.color.rgb = RGBColor(0x44, 0x72, 0xC4)
    run.font.size = Pt(10)
    run.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'AAAAAA')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)


# ══════════════════════════════════════════════════════════════════════════════
# PORTADA
# ══════════════════════════════════════════════════════════════════════════════
doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("UNIVERSIDAD TECNOLÓGICA DE VICTORIA\n")
r.bold = True; r.font.size = Pt(16); r.font.color.rgb = AZUL_TITULO

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = t2.add_run("Infraestructura en la Nube con Terraform\n")
r2.font.size = Pt(13); r2.font.color.rgb = AZUL_SUB

doc.add_paragraph()
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_r = title_p.add_run("Laboratorio 2\nDespliegue de Infraestructura Base en Azure\nResource Group & Virtual Network")
title_r.bold = True; title_r.font.size = Pt(22); title_r.font.color.rgb = AZUL_TITULO

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Proyecto: integradora  |  Ambiente: dev  |  Región: mexicocentral\n").font.size = Pt(11)
meta.add_run("Proveedor: azurerm ~> 5.0  |  Terraform >= 1.9.0").font.size = Pt(11)

doc.add_paragraph()
add_image_placeholder(doc, "Portada — logo institucional o captura del portal Azure")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 1. INTRODUCCIÓN
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "1. Introducción", 1)
add_hr(doc)
add_body(doc,
    "En este laboratorio se aplica el paradigma de Infraestructura como Código (IaC) utilizando "
    "Terraform para aprovisionar recursos base en Microsoft Azure. Los recursos creados son un "
    "Resource Group y una Virtual Network, los cuales constituyen la base de cualquier arquitectura "
    "en la nube empresarial.")
add_body(doc,
    "El laboratorio pone énfasis en buenas prácticas de Terraform: separación de archivos por "
    "responsabilidad, uso de variables con validación, locals para convenciones de nomenclatura, "
    "etiquetado consistente con tags y definición de outputs reutilizables.")

add_heading(doc, "Objetivos", 2, AZUL_SUB)
for obj in [
    "Configurar el proveedor azurerm con autenticación por subscription_id.",
    "Crear un Resource Group con nombre estandarizado mediante locals.",
    "Desplegar una Virtual Network con espacio de direcciones CIDR configurable.",
    "Aplicar etiquetas comunes a todos los recursos.",
    "Verificar los outputs generados tras el despliegue.",
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(obj)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 2. ESTRUCTURA DEL PROYECTO
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "2. Estructura del Proyecto", 1)
add_hr(doc)
add_body(doc,
    "El proyecto lab2 sigue la convención de separar cada responsabilidad en su propio archivo .tf, "
    "facilitando el mantenimiento y la legibilidad del código:")

add_code_block(doc,
"""lab2/
├── terraform.tf       # Versión de Terraform y proveedor requerido
├── providers.tf       # Configuración del proveedor azurerm
├── variables.tf       # Declaración de variables con tipos y validaciones
├── terraform.tfvars   # Valores concretos de las variables
├── locals.tf          # Lógica de nomenclatura y merge de tags
├── main.tf            # Recursos: Resource Group y Virtual Network
└── output.tf          # Outputs exportados del módulo""",
    label="Árbol de archivos — lab2/")

add_image_placeholder(doc, "Captura del explorador de archivos mostrando la carpeta lab2/")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 3. CÓDIGO FUENTE COMENTADO
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "3. Código Fuente", 1)
add_hr(doc)

# 3.1 terraform.tf
add_heading(doc, "3.1  terraform.tf — Bloque de configuración global", 2, AZUL_SUB)
add_body(doc,
    "Define la versión mínima de Terraform requerida y el proveedor hashicorp/azurerm en su rama 5.x. "
    "El constraint ~> 5.0 permite actualizaciones menores pero bloquea cambios mayores que rompan "
    "compatibilidad.")
add_code_block(doc,
"""terraform {
  required_version = ">= 1.9.0"

  required_providers {
    azurerm = {
      source  = "hashicorp/azurerm"
      version = "~> 5.0"
    }
  }
}""", label="terraform.tf")
add_image_placeholder(doc, "Captura de terraform.tf en el editor")

# 3.2 providers.tf
add_heading(doc, "3.2  providers.tf — Configuración del proveedor", 2, AZUL_SUB)
add_body(doc,
    "Configura el proveedor azurerm con el bloque features {} obligatorio. La subscription_id se "
    "lee desde la variable homónima marcada como sensitive para evitar su exposición en logs.")
add_code_block(doc,
"""provider "azurerm" {
  features {}
  subscription_id = var.subscription_id
}""", label="providers.tf")
add_image_placeholder(doc, "Captura de providers.tf en el editor")

# 3.3 variables.tf
add_heading(doc, "3.3  variables.tf — Declaración de variables", 2, AZUL_SUB)
add_body(doc,
    "Se definen seis variables con sus tipos, descripciones y, donde aplica, reglas de validación "
    "que se evalúan en tiempo de plan para detectar errores temprano.")
add_code_block(doc,
"""variable "project_name" {
  description = "Nombre corto del proyecto o aplicación."
  type        = string
  validation {
    condition     = length(var.project_name) >= 3 && length(var.project_name) <= 20
    error_message = "project_name debe contener entre 3 y 20 caracteres."
  }
}

variable "environment" {
  description = "Ambiente donde se desplegarán los recursos."
  type        = string
  validation {
    condition     = contains(["dev","test","qa","stg","prod"], var.environment)
    error_message = "environment debe ser dev, test, qa, stg o prod."
  }
}

variable "location" {
  description = "Región de Azure donde se desplegarán los recursos."
  type        = string
  default     = "mexicocentral"
}

variable "vnet_address_space" {
  description = "Espacio de direcciones CIDR asignado a la Virtual Network."
  type        = list(string)
  default     = ["10.0.0.0/16"]
}

variable "tags" {
  description = "Etiquetas comunes para los recursos de Azure."
  type        = map(string)
  default     = { managed_by = "terraform" }
}

variable "subscription_id" {
  description = "ID de la suscripción de Azure."
  type        = string
  sensitive   = true
}""", label="variables.tf")
add_image_placeholder(doc, "Captura de variables.tf en el editor")

# 3.4 terraform.tfvars
add_heading(doc, "3.4  terraform.tfvars — Valores de las variables", 2, AZUL_SUB)
add_body(doc,
    "Archivo que proporciona los valores reales para el ambiente de desarrollo del proyecto "
    "integradora. Terraform lo carga automáticamente por su nombre reservado.")
add_code_block(doc,
"""project_name = "integradora"
environment  = "dev"
location     = "mexicocentral"

vnet_address_space = ["10.10.0.0/16"]

tags = {
  managed_by  = "terraform"
  owner       = "it"
  cost_center = "integradora"
}""", label="terraform.tfvars")
add_image_placeholder(doc, "Captura de terraform.tfvars en el editor")

# 3.5 locals.tf
add_heading(doc, "3.5  locals.tf — Nomenclatura y tags consolidados", 2, AZUL_SUB)
add_body(doc,
    "Los locals centralizan la lógica de naming en un único lugar. La función merge() combina "
    "los tags del usuario con los atributos de proyecto, ambiente y región garantizando "
    "etiquetado consistente en todos los recursos.")
add_code_block(doc,
"""locals {
  resource_group_name  = "rg-${var.project_name}-${var.environment}-${var.location}-001"
  virtual_network_name = "vnet-${var.project_name}-${var.environment}-${var.location}-001"

  common_tags = merge(
    var.tags,
    {
      project     = var.project_name
      environment = var.environment
      location    = var.location
    }
  )
}""", label="locals.tf")
add_body(doc,
    "Ejemplo de nombres generados con los valores del tfvars:")
add_code_block(doc,
"""resource_group_name  → rg-integradora-dev-mexicocentral-001
virtual_network_name → vnet-integradora-dev-mexicocentral-001""",
    label="Nombres resultantes")
add_image_placeholder(doc, "Captura de locals.tf en el editor")

# 3.6 main.tf
add_heading(doc, "3.6  main.tf — Recursos de infraestructura", 2, AZUL_SUB)
add_body(doc,
    "Define los dos recursos principales. El Resource Group actúa como contenedor lógico en Azure "
    "y la Virtual Network se ancla a él para heredar región y grupo.")
add_code_block(doc,
"""resource "azurerm_resource_group" "main" {
  name     = local.resource_group_name
  location = var.location
  tags     = local.common_tags
}

resource "azurerm_virtual_network" "main" {
  name                = local.virtual_network_name
  location            = azurerm_resource_group.main.location
  resource_group_name = azurerm_resource_group.main.name
  address_space       = var.vnet_address_space
  tags                = local.common_tags
}""", label="main.tf")
add_image_placeholder(doc, "Captura de main.tf en el editor")

# 3.7 output.tf
add_heading(doc, "3.7  output.tf — Outputs del módulo", 2, AZUL_SUB)
add_body(doc,
    "Los outputs exponen atributos clave de los recursos creados. Son útiles para módulos raíz "
    "que consuman este módulo o para scripts de automatización post-despliegue.")
add_code_block(doc,
"""output "resource_group_name" {
  description = "Nombre del Resource Group creado"
  value       = azurerm_resource_group.main.name
}

output "resource_group_location" {
  description = "Ubicación del Resource Group"
  value       = azurerm_resource_group.main.location
}

output "virtual_network_name" {
  description = "Nombre de la Virtual Network creada"
  value       = azurerm_virtual_network.main.name
}

output "virtual_network_address_space" {
  description = "Espacio de direcciones de la Virtual Network"
  value       = azurerm_virtual_network.main.address_space
}

output "common_tags" {
  description = "Etiquetas comunes aplicadas a los recursos"
  value       = local.common_tags
}""", label="output.tf")
add_image_placeholder(doc, "Captura de output.tf en el editor")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 4. PRERREQUISITOS
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "4. Prerrequisitos", 1)
add_hr(doc)
add_body(doc, "Antes de ejecutar el laboratorio se deben cumplir los siguientes requisitos:")

prereqs = [
    ("Terraform >= 1.9.0",     "Descargar en https://developer.hashicorp.com/terraform/downloads"),
    ("Azure CLI >= 2.60",      "Descargar en https://learn.microsoft.com/cli/azure/install-azure-cli"),
    ("Cuenta Azure activa",    "Con una suscripción válida y permisos de Contributor"),
    ("subscription_id",        "ID de la suscripción configurado en terraform.tfvars o variable de entorno"),
    ("Editor de código",       "VS Code con extensión HashiCorp Terraform recomendado"),
]

tbl = doc.add_table(rows=1, cols=2)
tbl.style = 'Table Grid'
hdr = tbl.rows[0].cells
hdr[0].text = "Requisito"
hdr[1].text = "Detalle"
for cell in hdr:
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade_cell(cell, "1F497D")

for req, detail in prereqs:
    row = tbl.add_row().cells
    row[0].text = req
    row[1].text = detail

doc.add_paragraph()
add_image_placeholder(doc, "Captura: terraform version  y  az version  en terminal")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 5. FLUJO DE TRABAJO EN TERMINAL
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "5. Flujo de Trabajo en Terminal", 1)
add_hr(doc)
add_body(doc,
    "A continuación se describe cada paso del ciclo de vida de Terraform junto con los comandos "
    "exactos ejecutados en el laboratorio.")

# 5.1 Autenticación Azure
add_heading(doc, "5.1  Autenticación en Azure", 2, AZUL_SUB)
add_body(doc,
    "Se autentica la sesión de Azure CLI con la cuenta de trabajo. El comando abre una ventana del "
    "navegador para completar el inicio de sesión.")
add_code_block(doc,
"""# Iniciar sesión en Azure
az login

# Verificar la suscripción activa
az account show

# (Opcional) Cambiar a la suscripción del laboratorio
az account set --subscription "947b767e-d2d0-4136-938f-918c6fbc48c0\"""",
    label="Terminal — Autenticación Azure CLI")
add_image_placeholder(doc, "Captura del terminal: az login exitoso con cuenta mostrada")

# 5.2 terraform init
add_heading(doc, "5.2  Inicialización — terraform init", 2, AZUL_SUB)
add_body(doc,
    "Descarga el proveedor hashicorp/azurerm v5.8.0 y prepara el directorio de trabajo. "
    "El archivo .terraform.lock.hcl se genera o actualiza con los hashes del proveedor.")
add_code_block(doc,
"""# Posicionarse en el directorio del laboratorio
cd lab2

# Inicializar Terraform (descarga el proveedor azurerm)
terraform init""",
    label="Terminal — terraform init")
add_body(doc, "Salida esperada:")
add_code_block(doc,
"""Initializing the backend...
Initializing provider plugins...
- Finding hashicorp/azurerm versions matching "~> 5.0"...
- Installing hashicorp/azurerm v5.8.0...
- Installed hashicorp/azurerm v5.8.0 (signed by HashiCorp)

Terraform has been successfully initialized!""",
    label="Salida de terraform init")
add_image_placeholder(doc, "Captura del terminal: terraform init completado exitosamente")

# 5.3 terraform validate
add_heading(doc, "5.3  Validación — terraform validate", 2, AZUL_SUB)
add_body(doc,
    "Verifica la sintaxis y coherencia de todos los archivos .tf sin realizar ninguna llamada a la API de Azure.")
add_code_block(doc,
"""terraform validate""",
    label="Terminal — terraform validate")
add_code_block(doc,
"""Success! The configuration is valid.""",
    label="Salida esperada")
add_image_placeholder(doc, "Captura del terminal: terraform validate — Success!")

# 5.4 terraform fmt
add_heading(doc, "5.4  Formato — terraform fmt", 2, AZUL_SUB)
add_body(doc, "Formatea los archivos siguiendo el estilo canónico de HashiCorp.")
add_code_block(doc,
"""terraform fmt -recursive""",
    label="Terminal — terraform fmt")
add_image_placeholder(doc, "Captura del terminal: terraform fmt")

# 5.5 terraform plan
add_heading(doc, "5.5  Plan de ejecución — terraform plan", 2, AZUL_SUB)
add_body(doc,
    "Genera el plan de cambios sin modificar infraestructura. Se revisó que Terraform propone "
    "crear 2 recursos: el Resource Group y la Virtual Network.")
add_code_block(doc,
"""terraform plan""",
    label="Terminal — terraform plan")
add_body(doc, "Fragmento relevante de la salida:")
add_code_block(doc,
"""Terraform will perform the following actions:

  # azurerm_resource_group.main will be created
  + resource "azurerm_resource_group" "main" {
      + id       = (known after apply)
      + location = "mexicocentral"
      + name     = "rg-integradora-dev-mexicocentral-001"
      + tags     = {
          + "cost_center" = "integradora"
          + "environment" = "dev"
          + "location"    = "mexicocentral"
          + "managed_by"  = "terraform"
          + "owner"       = "it"
          + "project"     = "integradora"
        }
    }

  # azurerm_virtual_network.main will be created
  + resource "azurerm_virtual_network" "main" {
      + address_space = ["10.10.0.0/16"]
      + location      = "mexicocentral"
      + name          = "vnet-integradora-dev-mexicocentral-001"
      ...
    }

Plan: 2 to add, 0 to change, 0 to destroy.""",
    label="Salida de terraform plan")
add_image_placeholder(doc, "Captura completa del terminal: terraform plan — Plan: 2 to add")

# 5.6 terraform apply
add_heading(doc, "5.6  Aplicación — terraform apply", 2, AZUL_SUB)
add_body(doc,
    "Ejecuta los cambios del plan contra la API de Azure. Se confirma con 'yes' cuando Terraform "
    "solicita aprobación.")
add_code_block(doc,
"""terraform apply

# Para aprobación automática (entornos CI/CD):
terraform apply -auto-approve""",
    label="Terminal — terraform apply")
add_code_block(doc,
"""azurerm_resource_group.main: Creating...
azurerm_resource_group.main: Creation complete after 3s
azurerm_virtual_network.main: Creating...
azurerm_virtual_network.main: Creation complete after 8s

Apply complete! Resources: 2 added, 0 changed, 0 destroyed.

Outputs:

common_tags = {
  "cost_center" = "integradora"
  "environment" = "dev"
  "location"    = "mexicocentral"
  "managed_by"  = "terraform"
  "owner"       = "it"
  "project"     = "integradora"
}
resource_group_location          = "mexicocentral"
resource_group_name              = "rg-integradora-dev-mexicocentral-001"
virtual_network_address_space    = tolist(["10.10.0.0/16"])
virtual_network_name             = "vnet-integradora-dev-mexicocentral-001\"""",
    label="Salida de terraform apply")
add_image_placeholder(doc, "Captura del terminal: terraform apply — Apply complete!")
add_image_placeholder(doc, "Captura de los Outputs mostrados en terminal tras el apply")

# 5.7 terraform output
add_heading(doc, "5.7  Consulta de outputs — terraform output", 2, AZUL_SUB)
add_body(doc, "Permite consultar los outputs en cualquier momento sin necesidad de re-ejecutar apply.")
add_code_block(doc,
"""# Mostrar todos los outputs
terraform output

# Consultar un output específico
terraform output resource_group_name
terraform output virtual_network_name""",
    label="Terminal — terraform output")
add_image_placeholder(doc, "Captura del terminal: terraform output")

# 5.8 terraform show / state
add_heading(doc, "5.8  Inspección del estado", 2, AZUL_SUB)
add_code_block(doc,
"""# Ver el estado completo en formato legible
terraform show

# Listar recursos en el state
terraform state list

# Inspeccionar un recurso específico
terraform state show azurerm_resource_group.main
terraform state show azurerm_virtual_network.main""",
    label="Terminal — terraform show / state")
add_image_placeholder(doc, "Captura del terminal: terraform state list y state show")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 6. VERIFICACIÓN EN EL PORTAL DE AZURE
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "6. Verificación en el Portal de Azure", 1)
add_hr(doc)
add_body(doc,
    "Después del apply exitoso se verifica la creación de los recursos directamente en portal.azure.com.")

steps_azure = [
    ("Resource Group",
     "En el portal navegar a: Inicio → Grupos de recursos → rg-integradora-dev-mexicocentral-001\n"
     "Verificar: nombre, región (Mexico Central) y etiquetas."),
    ("Virtual Network",
     "Dentro del Resource Group hacer clic en: vnet-integradora-dev-mexicocentral-001\n"
     "Verificar: espacio de direcciones 10.10.0.0/16 y etiquetas."),
    ("Etiquetas",
     "En el recurso seleccionar la pestaña 'Etiquetas' y confirmar los 6 tags configurados:\n"
     "managed_by, owner, cost_center, project, environment, location."),
    ("Azure CLI (alternativa)",
     "Verificar por CLI sin abrir el portal:\n"
     "az group show --name rg-integradora-dev-mexicocentral-001\n"
     "az network vnet list --resource-group rg-integradora-dev-mexicocentral-001"),
]

for title, desc in steps_azure:
    add_heading(doc, f"6.x  {title}", 2, AZUL_SUB)
    add_body(doc, desc)
    add_image_placeholder(doc, f"Captura del Portal Azure — {title}")

# CLI verify block
add_code_block(doc,
"""# Verificar Resource Group desde CLI
az group show \\
  --name "rg-integradora-dev-mexicocentral-001" \\
  --output table

# Verificar Virtual Network desde CLI
az network vnet list \\
  --resource-group "rg-integradora-dev-mexicocentral-001" \\
  --output table

# Ver etiquetas del Resource Group
az group show \\
  --name "rg-integradora-dev-mexicocentral-001" \\
  --query tags""",
    label="Terminal — Verificación con Azure CLI")
add_image_placeholder(doc, "Captura del terminal: az group show y az network vnet list")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 7. DESTRUCCIÓN DE RECURSOS
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "7. Destrucción de Recursos", 1)
add_hr(doc)
add_body(doc,
    "Al finalizar el laboratorio se eliminan los recursos para evitar costos innecesarios en la suscripción.")
add_code_block(doc,
"""# Ver qué se destruirá antes de confirmar
terraform plan -destroy

# Destruir todos los recursos gestionados
terraform destroy

# Destrucción automática sin confirmación (CI/CD)
terraform destroy -auto-approve""",
    label="Terminal — terraform destroy")
add_code_block(doc,
"""azurerm_virtual_network.main: Destroying...
azurerm_virtual_network.main: Destruction complete after 5s
azurerm_resource_group.main: Destroying...
azurerm_resource_group.main: Destruction complete after 15s

Destroy complete! Resources: 2 destroyed.""",
    label="Salida esperada de terraform destroy")
add_image_placeholder(doc, "Captura del terminal: terraform destroy — Destroy complete!")
add_image_placeholder(doc, "Captura del Portal Azure confirmando que el Resource Group fue eliminado")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 8. RESUMEN DE RECURSOS CREADOS
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "8. Resumen de Recursos Creados", 1)
add_hr(doc)

tbl2 = doc.add_table(rows=1, cols=4)
tbl2.style = 'Table Grid'
hdrs = ["Recurso Terraform", "Tipo Azure", "Nombre Generado", "Región"]
for i, h in enumerate(hdrs):
    c = tbl2.rows[0].cells[i]
    c.text = h
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shade_cell(c, "1F497D")

rows_data = [
    ("azurerm_resource_group.main",   "Resource Group",   "rg-integradora-dev-mexicocentral-001",   "Mexico Central"),
    ("azurerm_virtual_network.main",  "Virtual Network",  "vnet-integradora-dev-mexicocentral-001", "Mexico Central"),
]
for rd in rows_data:
    r = tbl2.add_row().cells
    for i, v in enumerate(rd):
        r[i].text = v

doc.add_paragraph()
add_body(doc, "Tags aplicados a ambos recursos:")
add_code_block(doc,
"""managed_by  = "terraform"
owner       = "it"
cost_center = "integradora"
project     = "integradora"
environment = "dev"
location    = "mexicocentral\"""",
    label="Tags comunes")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 9. CONCLUSIONES
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "9. Conclusiones", 1)
add_hr(doc)
conclusiones = [
    "La separación de archivos .tf por responsabilidad mejora la mantenibilidad y permite "
    "escalar el proyecto sin modificar código no relacionado.",
    "El uso de variables con validaciones detecta errores de configuración en tiempo de plan, "
    "antes de realizar cambios en la infraestructura real.",
    "Los locals centralizan las convenciones de nomenclatura, evitando inconsistencias entre recursos.",
    "La función merge() garantiza que todos los recursos hereden los tags base más los específicos "
    "del proyecto, facilitando la gestión de costos y gobernanza en Azure.",
    "El ciclo init → validate → fmt → plan → apply representa una práctica estándar que minimiza "
    "riesgos en despliegues de infraestructura.",
    "Terraform genera un terraform.tfstate que actúa como fuente de verdad de la infraestructura, "
    "permitiendo detectar desvíos (drift) en futuras ejecuciones.",
]
for c in conclusiones:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(c)

doc.add_paragraph()
add_image_placeholder(doc, "Captura final — Portal Azure con ambos recursos creados y etiquetados")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
# 10. REFERENCIAS
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, "10. Referencias", 1)
add_hr(doc)
refs = [
    "HashiCorp. (2024). Terraform Documentation. https://developer.hashicorp.com/terraform/docs",
    "Microsoft. (2024). Azure Provider for Terraform. https://registry.terraform.io/providers/hashicorp/azurerm/latest/docs",
    "Microsoft. (2024). Azure Virtual Network Documentation. https://learn.microsoft.com/azure/virtual-network/",
    "Microsoft. (2024). Azure Resource Groups. https://learn.microsoft.com/azure/azure-resource-manager/management/manage-resource-groups-portal",
    "HashiCorp. (2024). Terraform azurerm_resource_group. https://registry.terraform.io/providers/hashicorp/azurerm/latest/docs/resources/resource_group",
    "HashiCorp. (2024). Terraform azurerm_virtual_network. https://registry.terraform.io/providers/hashicorp/azurerm/latest/docs/resources/virtual_network",
]
for ref in refs:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(ref).font.size = Pt(10)

# ── Guardar ───────────────────────────────────────────────────────────────────
out_path = r"c:\Users\HP\Documents\UTVT\Decimo\Almeida\terraform-main\terraform-main\lab2\Reporte_Lab2_Terraform_Azure.docx"
doc.save(out_path)
print(f"Documento guardado en:\n{out_path}")
